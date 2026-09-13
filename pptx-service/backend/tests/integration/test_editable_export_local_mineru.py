# SPDX-FileCopyrightText: 2026 北京京微资易科技有限公司 and AI WorkDeck contributors
# SPDX-License-Identifier: AGPL-3.0-or-later
"""
[checkba] 可编辑导出走本机 MinerU 引擎的集成用例（dev-board#617）。

这条用例是「真服务」用例：它需要本机 mineru-service 在跑，靠环境变量 MINERU_LOCAL_URL
指过去，不可达就 skip。**不要把它改成 mock**——本任务最容易假绿的地方正是「本地 /file_parse
返回的东西够不够提取器用」：提取器只认磁盘上的 layout.json（pdf_info[0] 的 para_blocks /
discarded_blocks 带 bbox）与 *_content_list.json，而本地通路以前只落 content_list，
于是元素表恒为空、导出静默退回纯图片。mock 掉解析服务就再也看不出这个缺口。

起法（与 desktop/main/services/mineru-service.js 同口径）：
  PACK=~/.aiworkdeck/packs/mineru-runtime/<version>
  MODELS=~/.aiworkdeck/models/mineru
  PYTHONPATH=$PACK/lib MINERU_DEVICE_MODE=cpu MINERU_MODEL_SOURCE=modelscope \
  MODELSCOPE_CACHE=$MODELS HF_HOME=$MODELS/hf MINERU_TOOLS_CONFIG_JSON=$MODELS/mineru.json \
  <bundled python3.11> -m mineru.cli.fast_api --host 127.0.0.1 --port <port>
然后 MINERU_LOCAL_URL=http://127.0.0.1:<port> 跑本文件。
"""
import os
import urllib.request

import pytest

LOCAL_URL = os.getenv('MINERU_LOCAL_URL', '').rstrip('/')


def _engine_reachable(url: str) -> bool:
    if not url:
        return False
    try:
        with urllib.request.urlopen(url + '/docs', timeout=5) as resp:
            return resp.status == 200
    except Exception:
        return False


pytestmark = [
    pytest.mark.integration,
    pytest.mark.requires_service,
    pytest.mark.slow,
    pytest.mark.skipif(
        not _engine_reachable(LOCAL_URL),
        reason='需要本机 MinerU 引擎在跑：设 MINERU_LOCAL_URL 指向 mineru-service',
    ),
]

FONT = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__)))), 'fonts', 'NotoSansSC-Regular.ttf')


def _render_slide(path: str) -> str:
    """画一张「标题 + 两段正文 + 2×3 表格」的幻灯片位图（不依赖任何模型/网络）。"""
    from PIL import Image, ImageDraw, ImageFont

    width, height = 1600, 900
    img = Image.new('RGB', (width, height), 'white')
    draw = ImageDraw.Draw(img)
    f_title = ImageFont.truetype(FONT, 64)
    f_body = ImageFont.truetype(FONT, 34)
    f_cell = ImageFont.truetype(FONT, 28)

    draw.text((120, 70), '季度经营情况汇报', font=f_title, fill='black')
    draw.text((120, 210), '第一段正文：本季度营业收入较上季度增长百分之十二。', font=f_body, fill='black')
    draw.text((120, 280), '第二段正文：主要来自法律服务与咨询两条业务线。', font=f_body, fill='black')

    rows = [['项目', '金额', '同比'], ['法律服务', '1200', '12%']]
    x0, y0, cell_w, cell_h = 200, 420, 380, 90
    for r in range(2):
        for c in range(3):
            x, y = x0 + c * cell_w, y0 + r * cell_h
            draw.rectangle([x, y, x + cell_w, y + cell_h], outline='black', width=3)
            draw.text((x + 24, y + 28), rows[r][c], font=f_cell, fill='black')

    img.save(path)
    return path


def _app(upload_folder: str):
    """最小 Flask 上下文：只提供这条链路真正读的那几个 config 键。"""
    from flask import Flask

    app = Flask(__name__)
    app.config.update(
        UPLOAD_FOLDER=upload_folder,
        MINERU_TOKEN='',            # 关键：全程无 token，正是发行版的样子
        MINERU_API_BASE='https://mineru.net',
        MINERU_LOCAL_URL=LOCAL_URL,
        MINERU_FORCE_CLOUD='0',
        # 背景重绘（inpaint）的 provider 在 from_defaults 里就会构造，构造时要有一把
        # AI key 才不抛。本用例只验版面提取，构造完就把重绘注册表清空，永不外呼。
        GOOGLE_API_KEY='not-used-in-this-test',
    )
    return app


def _analyze(tmp_path, png):
    """按生产口径（Flask config + from_defaults）分析一张图，返回 EditableImage。"""
    from services.image_editability import ImageEditabilityService, ServiceConfig
    from services.image_editability.inpaint_providers import InpaintProviderRegistry

    with _app(str(tmp_path)).app_context():
        config = ServiceConfig.from_defaults(
            extractor_method='mineru',   # 纯 MinerU：不碰百度高精度 OCR（外采）
            inpaint_method='generative',
            max_depth=1,
        )
        config.inpaint_registry = InpaintProviderRegistry()  # 不做背景重绘，不外呼图像模型
        return config, ImageEditabilityService(config).make_image_editable(png)


def test_editable_export_extracts_text_and_table_via_local_engine(monkeypatch, tmp_path):
    from services import file_parser_service

    png = _render_slide(str(tmp_path / 'slide.png'))

    # 记录所有出站 POST：既证明只打了本机引擎，也证明没有任何地方带 token
    posts = []
    real_post = file_parser_service.requests.post

    def spy_post(url, **kwargs):
        posts.append((url, kwargs.get('headers') or {}, kwargs.get('data') or {}))
        return real_post(url, **kwargs)

    monkeypatch.setattr(file_parser_service.requests, 'post', spy_post)

    # 无 token 也能构造（这正是修掉的那道硬门）
    config, result = _analyze(tmp_path, png)
    assert config.upload_folder == tmp_path
    extractor = config.extractor_registry.get_extractor(None)
    assert not extractor._parser_service.mineru_token, '本条用例必须在完全无 token 下跑通'

    types = [e.element_type for e in result.elements]
    texts = [e for e in result.elements
             if e.element_type in ('text', 'title') and (e.content or '').strip()]
    tables = [e for e in result.elements if e.element_type == 'table']

    assert len(texts) >= 3, f'标题+两段正文至少应提取到 3 个文本块，实际 {types}'
    assert tables, f'2×3 表格应被识别成 table 元素，实际 {types}'

    # 元素必须带真 bbox（提取器把 layout.json 的 bbox 缩放到位图尺寸）
    for element in result.elements:
        assert element.bbox.x1 > element.bbox.x0 and element.bbox.y1 > element.bbox.y0
        assert element.bbox.x1 <= result.width + 1 and element.bbox.y1 <= result.height + 1

    # layout.json 真落在提取器读的目录里（本次修复的核心：以前只落 content_list）
    layouts = list((tmp_path / 'mineru_files').glob('*/layout.json'))
    assert layouts, 'layout.json 必须落在 UPLOAD_FOLDER/mineru_files/<extract_id>/ 下'
    content_lists = list((tmp_path / 'mineru_files').glob('*/*_content_list.json'))
    assert content_lists, '提取器要求 content_list 与 layout.json 同目录共存'

    assert posts, '没有任何出站请求说明根本没调到解析服务'
    for url, headers, data in posts:
        assert url.startswith(LOCAL_URL), f'出现了非本机引擎的请求：{url}'
        assert 'Authorization' not in headers and 'authorization' not in headers, \
            f'本地通路不该带鉴权头：{headers}'
        assert not any('token' in str(k).lower() for k in data), f'表单里出现了 token 字段：{data}'


def test_analyzed_slide_assembles_into_pptx_with_editable_text(tmp_path):
    """
    提取到的元素要真能装配成 PPTX：正文/标题落成可编辑文本框。

    表格在纯本机通路下以图片块放回（可编辑单元格要靠递归 + 重绘背景，依赖外部 OCR/重绘
    能力），所以这里只对文字断言「可编辑」，对表格断言「有图形落位」——这也是工具描述里
    对用户的口径，别在这条断言上写超出实现的承诺。
    """
    from pptx import Presentation
    from services.export_service import ExportService

    png = _render_slide(str(tmp_path / 'slide.png'))
    _, analyzed = _analyze(tmp_path, png)

    out = str(tmp_path / 'editable.pptx')
    with _app(str(tmp_path)).app_context():
        ExportService.create_editable_pptx_with_recursive_analysis(
            editable_images=[analyzed], output_file=out,
            max_depth=1, max_workers=1, fail_fast=False)

    slides = list(Presentation(out).slides)
    assert len(slides) == 1
    texts = [sh.text_frame.text.strip() for sh in slides[0].shapes
             if sh.has_text_frame and sh.text_frame.text.strip()]
    assert len(texts) >= 3, f'标题与两段正文都应落成可编辑文本框，实际 {texts}'
    assert any('季度经营情况汇报' in t for t in texts), f'标题文字丢了：{texts}'
    pictures = [sh for sh in slides[0].shapes if sh.shape_type == 13]
    assert pictures, '表格区域应有图形落位（当前实现以图片块保留）'
